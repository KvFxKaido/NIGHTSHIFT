"""Export the Midnight Club Reference DB workbook to text.

    python design/reference/midnight-club/export.py <workbook.xlsx>

Writes data/*.csv (one per Excel table, values not formulas) and digest.md
(the narrative tables and summary sheets as markdown) beside this file.
README.md is written by hand and is not touched. Needs openpyxl.
"""

import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries

HERE = Path(__file__).resolve().parent

# Excel table name -> csv name. Every table in the workbook must be listed, so a
# new table in a later version fails loudly instead of being silently dropped.
TABLES = {
    "VehiclesTable": "mc3-vehicles",
    "UnlocksTable": "mc3-unlocks",
    "ProgressionTable": "mc3-progression",
    "NightShiftNotesTable": "mc3-nightshift-notes",
    "MCLAVehicles": "mcla-vehicles",
    "MCLAProgression": "mcla-ranks",
    "MCLAOtherUnlocks": "mcla-other-unlocks",
    "MCLARewards": "mcla-rewards",
    "MCLATuneGroups": "mcla-tune-groups",
    "MCLAEconomyObservations": "mcla-economy",
    "MCLAREPPayouts": "mcla-rep-payouts",
    "MCLADesignNotes": "mcla-design-notes",
    "MCLASources": "mcla-sources",
    "MCLAConflicts": "mcla-conflicts",
}

# Tables short or wordy enough to read as markdown as well as grep as csv.
DIGEST_TABLES = [
    ("MC3 progression", "ProgressionTable"),
    ("MC3 patterns for NightShift", "NightShiftNotesTable"),
    ("MCLA career ranks", "MCLAProgression"),
    ("MCLA other unlock tracks", "MCLAOtherUnlocks"),
    ("MCLA tune groups", "MCLATuneGroups"),
    ("MCLA economy observations", "MCLAEconomyObservations"),
    ("MCLA prize vehicles", "MCLARewards"),
    ("MCLA patterns for NightShift", "MCLADesignNotes"),
    ("MCLA source conflicts", "MCLAConflicts"),
]

# Columns holding MC3 source URLs, rewritten to the workbook's S-ids.
SOURCE_COLUMNS = {"Stats Source", "Unlock Source", "Source", "Evidence source"}

# Summary sheets have no Excel tables; their blocks are named by hand.
# (sheet, title, range, first row is a header)
REGIONS = [
    ("Dashboard", "MC3 roster totals", "A5:B10", False),
    ("Dashboard", "MC3 by class", "D5:H10", True),
    ("Dashboard", "MC3 by category", "A13:C21", True),
    ("Dashboard", "MC3 key progression facts", "D14:E21", False),
    ("Cross-game", "Roster metrics", "A5:D12", True),
    ("Cross-game", "Category mapping", "A15:D24", True),
    ("Cross-game", "MCLA purchase groups", "A27:D31", True),
    ("Cross-game", "Coverage and limits", "A34:D38", True),
    ("Sources", "MC3 sources", "A1:F8", True),
    ("Sources", "MC3 methodology and caveats", "A12:B17", False),
    ("MCLA Sources", "MCLA methodology and caveats", "A41:B50", False),
]

# Free-standing notes on summary sheets, quoted in the digest.
NOTES = [("Dashboard", "A24"), ("Cross-game", "A41"), ("Cross-game", "A43"), ("Cross-game", "A45")]


def text(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.3f}".rstrip("0").rstrip(".")
    return str(value).strip()


def grid(ws, ref):
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return [
        [text(ws.cell(r, c).value) for c in range(min_col, max_col + 1)]
        for r in range(min_row, max_row + 1)
    ]


def source_ids(wb):
    """MC3 source URL -> its S-id, so vehicle rows cite S1 rather than a URL."""
    ids = {}
    for row in grid(wb["Sources"], "A2:F8"):
        if row[0] and row[4]:
            ids[row[4]] = row[0]
    return ids


def cite(cell, ids):
    parts = [p.strip() for p in cell.split("|")]
    if all(p in ids for p in parts):
        return "; ".join(ids[p] for p in parts)
    return cell


def md_table(rows, header):
    rows = [[c.replace("|", "\\|").replace("\n", " ") for c in r] for r in rows]
    if not header:
        rows = [[""] * len(rows[0])] + rows
    widths = len(rows[0])
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * widths]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ids = source_ids(wb)
    tables = {}
    for ws in wb.worksheets:
        for name, ref in ws.tables.items():
            if name not in TABLES:
                sys.exit(f"table {name} on {ws.title} has no csv name in TABLES")
            tables[name] = (ws, ref)
    missing = set(TABLES) - set(tables)
    if missing:
        sys.exit(f"workbook no longer has tables: {sorted(missing)}")

    data = HERE / "data"
    data.mkdir(exist_ok=True)
    exported = {}
    for name, (ws, ref) in tables.items():
        rows = grid(ws, ref)
        for col, title in enumerate(rows[0]):
            if title in SOURCE_COLUMNS:
                for row in rows[1:]:
                    row[col] = cite(row[col], ids)
        with open(data / f"{TABLES[name]}.csv", "w", newline="", encoding="utf-8") as f:
            csv.writer(f, lineterminator="\n").writerows(rows)
        exported[name] = rows

    title = text(wb["Cross-game"]["A1"].value)
    lines = [
        "# Midnight Club reference digest",
        "",
        f"Generated by `export.py` from `{Path(path).name}`",
        f"({title}). Do not edit by hand; change the workbook and export again.",
        "Blank means not verified or not applicable, never zero. The full tables",
        "are in `data/`; see `README.md`.",
        "",
        "The Dashboard and MC3 tabs are the workbook's preserved v0.1 snapshot; their",
        "own title still says v0.1.",
    ]
    for sheet, heading, ref, header in REGIONS:
        lines += ["", f"## {heading}", "", f"*Sheet: {sheet}, {ref}.*", ""]
        lines.append(md_table(grid(wb[sheet], ref), header))
    lines += ["", "## Notes on the summary sheets", ""]
    lines += [f"- {text(wb[s][c].value)} *({s})*" for s, c in NOTES]
    for heading, name in DIGEST_TABLES:
        rows = exported[name]
        lines += ["", f"## {heading}", "", f"*`data/{TABLES[name]}.csv`, {len(rows) - 1} rows.*", ""]
        lines.append(md_table(rows, True))
    (HERE / "digest.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    for name, rows in exported.items():
        print(f"{TABLES[name]}.csv: {len(rows) - 1} rows")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
